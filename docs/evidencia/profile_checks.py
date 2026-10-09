import os
from pathlib import Path
SALIDA = Path(__file__).resolve().parent
import pandas as pd, numpy as np, re, unicodedata, zipfile
pd.set_option('display.width', 250); pd.set_option('display.max_rows', 300)
base = str(Path(os.environ.get("ANID_CRUDO") or (Path(__file__).resolve().parents[2] / "data" / "crudo")) / "")
out = open(SALIDA/'profile_checks.txt','w',encoding='utf-8')
def P(*a):
    s=' '.join(str(x) for x in a); print(s); out.write(s+'\n')

h = pd.read_csv(base+'Historico-de-Proyectos-Adjudicados/BDH_HISTORICA.csv', dtype=str, keep_default_na=False, sep=';', encoding='utf-8-sig')
h['anno'] = pd.to_numeric(h['AGNO_FALLO'], errors='coerce')
h['monto'] = pd.to_numeric(h['MONTO_ADJUDICADO'].str.replace('.','',regex=False).str.replace(',','.',regex=False), errors='coerce')

P('### A) filas MONEDA="M$" (94)')
s = h[h['MONEDA']=='M$']
P('  años:', s['anno'].value_counts().sort_index().to_dict())
P('  montos min/median/max:', s['monto'].min(), s['monto'].median(), s['monto'].max())
P('  ejemplos:', s[['CODIGO_PROYECTO','AGNO_FALLO','INSTRUMENTO','monto']].head(6).to_string(index=False))
P('  concursos:', s['NOMBRE_CONCURSO'].value_counts().head(5).to_dict())

P('\n### B) AGNO_FALLO >= 2026 (corte declarado 31-12-2025)')
f = h[h['anno']>=2026]
P('  filas:', len(f), '| años:', f['anno'].value_counts().sort_index().to_dict())
P('  subdireccion:', f['SUBDIRECCION'].value_counts().to_dict())
P('  programa:', f['PROGRAMA'].value_counts().to_dict())
P('  concursos top:', f['NOMBRE_CONCURSO'].value_counts().head(6).to_dict())
P('  regiones:', f['REGION_EJECUCION'].value_counts().head(8).to_dict())
P('  nulos monto:', f['monto'].isna().sum())

P('\n### C) region sin dato (global)')
P('  SIN INFORMACION+EXTRANJERO+MULTIREGIONAL:', int(h['REGION_EJECUCION'].isin(['SIN INFORMACION','EXTRANJERO','MULTIREGIONAL']).sum()), 'de', len(h))
P('  filas 2016+ sin region asignable:', int(h[(h['anno']>=2016)]['REGION_EJECUCION'].isin(['SIN INFORMACION','EXTRANJERO']).sum()))

P('\n### D) institucion x region: la jerarquia no es 1:1?')
hh = h[h['REGION_EJECUCION'].str.startswith('REGION')]
g = hh.groupby('INSTITUCION_PRINCIPAL')['REGION_EJECUCION'].nunique().sort_values(ascending=False)
P('  instituciones (en regiones reales):', len(g), '| con >1 region:', int((g>1).sum()), '| con >2:', int((g>2).sum()))
for k,v in g.head(6).items():
    P('   %-50s %d regiones: %s' % (k, v, hh.loc[hh['INSTITUCION_PRINCIPAL']==k,'REGION_EJECUCION'].value_counts().head(4).to_dict()))

P('\n### E) instituciones con volumen suficiente para filtros')
vc = h['INSTITUCION_PRINCIPAL'].value_counts()
P('  >=100 proyectos:', int((vc>=100).sum()), '| >=30:', int((vc>=30).sum()), '| >=10:', int((vc>=10).sum()), '| >=1:', len(vc))
P('  cola larga: instituciones con <=3 proyectos:', int((vc<=3).sum()), '=', round((vc<=3).mean()*100,1), '% de las instituciones y', int(vc[vc<=3].sum()), 'filas =', round(vc[vc<=3].sum()/len(h)*100,1),'% del total')

P('\n### F) ejemplo de cierre: monto adjudicado por MACROZONA, 2020-2025, solo M$')
w = h[(h['anno'].between(2020,2025)) & (h['MONEDA'].str.contains('pesos', case=False)) & h['monto'].notna()].copy()
t = w.groupby('MACROZONA_MINCIENCIA')['monto'].agg(filas='count', total_M_miles='sum')
t['total_millones_CLP'] = (t['total_M_miles']*1000/1e6).round(0)
P(t.sort_values('total_M_miles', ascending=False).to_string())
P('  total filas usadas:', len(w), 'de', int(h['anno'].between(2020,2025).sum()), 'del periodo')

P('\n### G) top 10 instituciones 2020-2025 (n y monto M$)')
t2 = w.groupby('INSTITUCION_PRINCIPAL')['monto'].agg(n='count', monto_M='sum').sort_values('n', ascending=False).head(10)
P(t2.round(0).to_string())

P('\n### H) xlsx vs csv (misma data?)')
for name, xl, csv in [('Historico','Historico-de-Proyectos-Adjudicados/BDH_HISTORICA.xlsx','Historico-de-Proyectos-Adjudicados/BDH_HISTORICA.csv'),
                      ('Colaboracion','Colaboracion_Institucional/CONSOLIDADO_INSTITUCIONES.xlsx','Colaboracion_Institucional/CONSOLIDADO_INSTITUCIONES.csv'),
                      ('Postulaciones','Postulaciones/POSTULACIONES_2016_2025.xlsx','Postulaciones/POSTULACIONES_2016_2025.csv'),
                      ('Becas','Becas/BDH_BECAS SCH.xlsx','Becas/BDH_BECAS SCH.csv')]:
    import openpyxl
    wb = openpyxl.load_workbook(base+xl, read_only=True)
    ws = wb[wb.sheetnames[0]]
    P('  %-14s xlsx sheet=%r dims=%sx%s | csv lines=%d' % (name, wb.sheetnames[0], ws.max_row, ws.max_column, sum(1 for _ in open(base+csv, encoding='utf-8-sig'))))
    wb.close()

P('\n### I) licencias')
for r in ['Historico-de-Proyectos-Adjudicados','Colaboracion_Institucional','Postulaciones','Becas','Palabras_claves','Instructivo_descarga','FONDECYT-Abstracts','Consulta-Publica-Politica-de-Acesso-Abierto']:
    import subprocess
    o = subprocess.run(['gh','api','repos/ANID-GITHUB/'+r,'--jq','(.license.spdx_id // "SIN LICENCIA") + " | " + (.size|tostring) + " KB | " + .pushed_at'], capture_output=True, text=True).stdout.strip()
    P('  %-45s %s' % (r, o))
out.close()
